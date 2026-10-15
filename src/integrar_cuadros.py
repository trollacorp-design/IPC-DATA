"""
Fase 4 - Integra las tablas oficiales del BCRD (Cuadros 3.1, 3.3 y 3.8, y el Anexo 3).

Entradas (en data/raw, se buscan por patron de nombre):
  Cuadros*2025*.xlsx   -> hojas CUADRO 3.1 (IPC general), 3.3 (12 grupos), 3.8 (subyacente)
  Anexos*.xlsx         -> hoja Anexo 3 (clasificacion bien/servicio, transable, subyacente)
  data/processed/ipc_largo.parquet (fase 2), que aporta 2026 y la jerarquia de Salud

Salidas (data/processed):
  ipc_series_oficiales.parquet   general + 12 grupos, 1999-2026, con variaciones
  subyacente.parquet             IPC general y subyacente, 2000-2026
  anexo3_clasificacion.parquet   364 articulos con su clasificacion
  salud_componentes.parquet      los 49 nodos de Salud con peso y clasificacion

Como correrlo, desde la raiz del proyecto:  python src/integrar_cuadros.py
"""
import re
import unicodedata
from pathlib import Path

import openpyxl
import pandas as pd

RAW = Path("data/raw")
OUT = Path("data/processed")
MESES = {m: i + 1 for i, m in enumerate(
    ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
     "agosto", "septiembre", "octubre", "noviembre", "diciembre"])}
GRUPOS = [f"{i:02d}" for i in range(1, 13)]

# Subyacente interanual 2026: la tabla oficial (Cuadro 3.8) llega a dic 2025, asi que estos tres
# valores se toman de los informes mensuales del BCRD (Informe-IPC-2026-01/02/03.pdf)
SUBYACENTE_2026 = {"2026-01-01": 4.89, "2026-02-01": 4.76, "2026-03-01": 4.58}


def buscar(patron: str) -> Path:
    hallados = sorted(RAW.glob(patron))
    if not hallados:
        raise FileNotFoundError(f"No encuentro '{patron}' en {RAW.resolve()}")
    return hallados[0]


def norm(s) -> str:
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]", "", s)).strip()


def num(v):
    return pd.to_numeric(v, errors="coerce")


# ----------------------------------------------------------------------------- lectura de cuadros
def leer_general(wb) -> pd.DataFrame:
    """Cuadro 3.1: indice general (base anual) y sus variaciones oficiales."""
    ws, filas, anio = wb["CUADRO 3.1"], [], None
    for r in range(8, ws.max_row + 1):
        a, m = ws.cell(r, 1).value, ws.cell(r, 2).value
        if a is not None and str(a).strip().isdigit():
            anio = int(str(a).strip())
        if isinstance(m, str) and m.strip().lower() in MESES and anio:
            filas.append({"fecha": pd.Timestamp(anio, MESES[m.strip().lower()], 1),
                          "indice": num(ws.cell(r, 5).value), "of_mensual": num(ws.cell(r, 6).value),
                          "of_acumulada": num(ws.cell(r, 7).value), "of_interanual": num(ws.cell(r, 8).value)})
    return pd.DataFrame(filas).set_index("fecha")


def leer_grupos(wb) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Cuadro 3.3: indice y variacion mensual de los 12 grupos, mas el bloque anual (dic a dic).
    Los grupos van siempre en el mismo orden (01 a 12), asi que el codigo se asigna por posicion."""
    ws = wb["CUADRO 3.3"]
    cols = {c: (f"{i + 1:02d}", str(ws.cell(5, c).value).strip()) for i, c in enumerate(range(2, 25, 2))}
    mensual, anual, anio = [], [], None
    for r in range(8, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, str) and a.strip().lower().startswith("dic."):          # bloque anual
            for c, (cod, n) in cols.items():
                anual.append({"fecha": pd.Timestamp(int(a.strip()[-4:]), 12, 1), "codigo": cod,
                              "of_anual": num(ws.cell(r, c + 1).value)})
        elif a is not None and str(a).strip().isdigit():                          # encabezado de ano
            anio = int(str(a).strip())
        elif isinstance(a, str) and a.strip().lower() in MESES and anio:          # fila de mes
            f = pd.Timestamp(anio, MESES[a.strip().lower()], 1)
            for c, (cod, n) in cols.items():
                mensual.append({"fecha": f, "codigo": cod, "grupo": n, "indice": num(ws.cell(r, c).value),
                                "of_mensual": num(ws.cell(r, c + 1).value)})
    return pd.DataFrame(mensual), pd.DataFrame(anual)


def leer_subyacente(wb) -> pd.DataFrame:
    """Cuadro 3.8: IPC general y subyacente con sus variaciones."""
    ws, filas, anio = wb["CUADRO 3.8"], [], None
    for r in range(7, ws.max_row + 1):
        a, m = ws.cell(r, 1).value, ws.cell(r, 2).value
        if a is not None and str(a).strip().isdigit():
            anio = int(str(a).strip())
        if isinstance(m, str) and m.strip().lower() in MESES and anio:
            filas.append({"fecha": pd.Timestamp(anio, MESES[m.strip().lower()], 1),
                          "ipc_general": num(ws.cell(r, 3).value), "gen_interanual": num(ws.cell(r, 6).value),
                          "ipc_subyacente": num(ws.cell(r, 7).value), "suby_mensual": num(ws.cell(r, 8).value),
                          "suby_acumulada": num(ws.cell(r, 9).value), "suby_interanual": num(ws.cell(r, 10).value)})
    return pd.DataFrame(filas).sort_values("fecha").reset_index(drop=True)


# ----------------------------------------------------------------------------- armado de series
def armar_series(gen, grp_m, grp_a, largo) -> pd.DataFrame:
    nombres = (largo[largo["nivel"].isin(["grupo", "general"])]
               .drop_duplicates("codigo").set_index("codigo")["nombre"])
    # Control: el nombre de cada grupo en el Cuadro 3.3 empieza igual que en el Excel de ONE
    for cod, g in grp_m.drop_duplicates("codigo")[["codigo", "grupo"]].itertuples(index=False):
        assert norm(g).split()[0] == norm(nombres[cod]).split()[0], f"Grupo {cod}: '{g}' vs '{nombres[cod]}'"

    base = pd.concat([
        gen.reset_index().assign(codigo="GEN")[["fecha", "codigo", "indice"]],
        grp_m[["fecha", "codigo", "indice"]]])
    base = base[(base["fecha"] >= "1999-01-01") & (base["fecha"] <= "2025-12-01")]

    # 2026: del Excel de ONE (fase 2). Se comprueba que ONE y los Cuadros son la misma serie.
    one = largo[largo["codigo"].isin(["GEN"] + GRUPOS)][["fecha", "codigo", "indice"]]
    sol = base.merge(one, on=["fecha", "codigo"], suffixes=("", "_one"))
    assert len(sol) == 13 * 63, f"Solape inesperado: {len(sol)} filas"
    assert (sol["indice"] - sol["indice_one"]).abs().max() < 0.002, "ONE y Cuadros no coinciden"
    base = pd.concat([base, one[one["fecha"] >= "2026-01-01"]])

    ancho = base.pivot(index="fecha", columns="codigo", values="indice").sort_index()
    assert ancho.notna().all().all(), "Hay meses sin dato"
    dic = ancho[ancho.index.month == 12].copy()
    dic.index = dic.index.year + 1
    ref = dic.reindex(ancho.index.year)
    ref.index = ancho.index
    tablas = {"indice": ancho, "var_mensual": ancho.pct_change() * 100,
              "var_acumulada": (ancho / ref - 1) * 100, "var_interanual": ancho.pct_change(12) * 100}
    largo_s = None
    for nombre, t in tablas.items():
        plano = t.reset_index().melt(id_vars="fecha", var_name="codigo", value_name=nombre)
        largo_s = plano if largo_s is None else largo_s.merge(plano, on=["fecha", "codigo"])
    largo_s["nombre"] = largo_s["codigo"].map(nombres)
    largo_s["en_ventana"] = (largo_s["fecha"] >= "2021-01-01") & (largo_s["fecha"] <= "2026-03-01")
    return largo_s.sort_values(["codigo", "fecha"]).reset_index(drop=True)


def validar_series(s, gen, grp_m, grp_a) -> None:
    ancho = lambda col: s.pivot(index="fecha", columns="codigo", values=col)
    # 1) General: variaciones calculadas = oficiales del Cuadro 3.1 (1999-2025)
    g = s[s["codigo"] == "GEN"].set_index("fecha")
    comun = g.index.intersection(gen.index)
    for mio, of in [("var_mensual", "of_mensual"), ("var_interanual", "of_interanual")]:
        d = (g.loc[comun, mio] - gen.loc[comun, of]).dropna().abs().max()
        assert d < 0.005, f"{mio} general no coincide con el Cuadro 3.1 ({d:.4f})"
    # 2) Grupos: variacion mensual calculada = columna Var.(%) del Cuadro 3.3
    calc_m = s.set_index(["fecha", "codigo"])["var_mensual"]
    of_m = grp_m.dropna(subset=["of_mensual"]).set_index(["fecha", "codigo"])["of_mensual"]
    comun = of_m.index.intersection(calc_m.dropna().index)
    d = (of_m[comun] - calc_m[comun]).abs().max()
    assert len(comun) > 3500 and d < 0.005, f"Variacion mensual de grupos no coincide ({d:.4f}, n={len(comun)})"
    # 3) Grupos: variacion dic a dic calculada = bloque anual del Cuadro 3.3.
    #    El bloque anual publicado tiene 6 celdas (2013-2017) que no cuadran ni con los indices ni con
    #    las variaciones mensuales de la misma hoja. Se exige coincidencia exacta desde 2018 y, en las
    #    celdas dudosas, se comprueba que los indices SI cuadran con las variaciones mensuales.
    calc_a = s[s["fecha"].dt.month == 12].set_index(["fecha", "codigo"])["var_interanual"]
    of_a = grp_a.dropna(subset=["of_anual"]).set_index(["fecha", "codigo"])["of_anual"]
    comun = of_a.index.intersection(calc_a.dropna().index)
    dif = (of_a[comun] - calc_a[comun]).abs()
    malas = dif[dif > 0.005]
    assert len(comun) > 300, "Muy pocas variaciones anuales para comparar"
    assert all(f.year < 2018 for f, _ in malas.index), "Hay diferencias en el bloque anual desde 2018"
    for (f, c) in malas.index:
        mm = grp_m[(grp_m["codigo"] == c) & (grp_m["fecha"].dt.year == f.year)]["of_mensual"]
        compuesta = ((1 + mm / 100).prod() - 1) * 100
        assert abs(compuesta - calc_a[(f, c)]) < 0.005, f"Indices y variaciones mensuales no cuadran ({f:%Y} grupo {c})"
    print(f"Nota: {len(malas)} de {len(comun)} celdas del bloque anual del Cuadro 3.3 (2013-2017) no cuadran "
          "con la propia tabla; se usan los indices.")

    # 4) Informe de enero 2021 (serie empalmada oficial): variacion interanual, 12 grupos y general
    ene21 = [8.87, 6.41, -0.47, 6.29, 5.68, 6.45, 9.36, 0.67, 3.59, -4.86, 5.07, 5.57]
    v = ancho("var_interanual").loc["2021-01-01"]
    for c, of in zip(GRUPOS, ene21):
        assert abs(v[c] - of) < 0.006, f"Interanual ene-2021 grupo {c}: {v[c]:.3f} vs {of}"
    assert abs(v["GEN"] - 6.23) < 0.006
    # 5) Informe de marzo 2026: variacion interanual, 12 grupos y general
    mar26 = [6.90, 5.16, -1.28, 2.62, 1.81, 5.53, 3.13, -1.18, 1.66, 7.15, 7.41, 6.01]
    v = ancho("var_interanual").loc["2026-03-01"]
    for c, of in zip(GRUPOS, mar26):
        assert abs(v[c] - of) < 0.006, f"Interanual mar-2026 grupo {c}: {v[c]:.3f} vs {of}"
    assert abs(v["GEN"] - 4.63) < 0.006
    print("Validaciones de series OK")


def armar_subyacente(wb) -> pd.DataFrame:
    sub = leer_subyacente(wb)
    # Dic 2025 oficial: general 4.95 y subyacente 4.85 (grafica del informe de marzo 2026)
    d = sub[sub["fecha"] == "2025-12-01"].iloc[0]
    assert abs(d["gen_interanual"] - 4.95) < 0.006 and abs(d["suby_interanual"] - 4.85) < 0.006
    extra = pd.DataFrame({"fecha": pd.to_datetime(list(SUBYACENTE_2026)),
                          "suby_interanual": list(SUBYACENTE_2026.values())})
    sub["fuente"] = "Cuadro 3.8 BCRD"
    extra["fuente"] = "Informe mensual BCRD 2026"
    return pd.concat([sub, extra], ignore_index=True).sort_values("fecha").reset_index(drop=True)


# ----------------------------------------------------------------------------- Anexo 3
def leer_anexo3(path: Path, largo: pd.DataFrame) -> pd.DataFrame:
    ws = openpyxl.load_workbook(path, data_only=True)["Anexo 3 Clasif. Subíndices"]
    filas = []
    for r in range(7, ws.max_row + 1):
        g, a, p = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value
        if g is not None and str(g).strip().isdigit() and a and p is not None:
            filas.append({"grupo_cod": f"{int(str(g).strip()):02d}", "articulo": str(a).strip(),
                          "pond_anexo": float(p), "bs": ws.cell(r, 4).value,
                          "tnt": ws.cell(r, 5).value, "suby": ws.cell(r, 6).value})
    a3 = pd.DataFrame(filas)
    a3["tipo"] = a3["bs"].map({"B": "Bien", "S": "Servicio"})
    a3["transable"] = a3["tnt"].map({"T": "Transable", "NT": "No transable"})
    a3["subyacente"] = a3["suby"].map({"Suby": "Subyacente", "Ex": "Excluido"})
    assert a3[["tipo", "transable", "subyacente"]].notna().all().all(), "Clasificacion con valores inesperados"

    art = (largo[(largo["nivel"] == "articulo") & (largo["fecha"] == largo["fecha"].min())]
           [["codigo", "nombre", "grupo_cod", "ponderacion"]])
    a3["k"], art["k"] = a3["articulo"].map(norm), art["nombre"].map(norm)
    a3["p"], art["p"] = a3["pond_anexo"].round(6), art["ponderacion"].round(6)
    # 1.a pasada: grupo + nombre. 2.a pasada (4 articulos con nombre distinto): grupo + ponderacion
    m1 = a3.merge(art[["codigo", "grupo_cod", "k"]], on=["grupo_cod", "k"], how="left")
    falta = m1["codigo"].isna()
    m2 = a3[falta.values].merge(art[["codigo", "grupo_cod", "p"]], on=["grupo_cod", "p"], how="left")
    m1.loc[falta.values, "codigo"] = m2["codigo"].values
    assert m1["codigo"].notna().all() and m1["codigo"].is_unique, "No se pudieron emparejar todos los articulos"
    assert len(m1) == 364

    # Controles contra el documento metodologico y el Cuadro 3.8
    assert abs(m1.loc[m1["subyacente"] == "Excluido", "pond_anexo"].sum() - 30.19) < 0.01, "Peso excluido != 30.19%"
    salud = m1[m1["grupo_cod"] == "06"]
    assert len(salud) == 32 and abs(salud["pond_anexo"].sum() - 4.74) < 0.01
    return m1[["codigo", "grupo_cod", "articulo", "pond_anexo", "tipo", "transable", "subyacente"]]


def armar_salud(largo: pd.DataFrame, a3: pd.DataFrame) -> pd.DataFrame:
    nodos = (largo[(largo["grupo_cod"] == "06") & (largo["fecha"] == largo["fecha"].min())]
             [["codigo", "nombre", "nivel", "grupo_cod", "subgrupo_cod", "clase_cod", "subclase_cod", "ponderacion"]])
    nodos = nodos.merge(a3[["codigo", "tipo", "transable", "subyacente"]], on="codigo", how="left")
    peso_salud = nodos.loc[nodos["nivel"] == "grupo", "ponderacion"].iloc[0]
    nodos["parte_de_salud_pct"] = nodos["ponderacion"] / peso_salud * 100
    assert len(nodos) == 49
    return nodos.sort_values("codigo").reset_index(drop=True)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    largo = pd.read_parquet(OUT / "ipc_largo.parquet")
    wb = openpyxl.load_workbook(buscar("Cuadros*2025*.xlsx"), data_only=True)

    gen = leer_general(wb)
    grp_m, grp_a = leer_grupos(wb)
    series = armar_series(gen, grp_m, grp_a, largo)
    validar_series(series, gen, grp_m, grp_a)
    sub = armar_subyacente(wb)
    a3 = leer_anexo3(buscar("Anexos*.xlsx"), largo)
    salud = armar_salud(largo, a3)
    print("Validaciones del Anexo 3 OK")

    series.to_parquet(OUT / "ipc_series_oficiales.parquet", index=False)
    sub.to_parquet(OUT / "subyacente.parquet", index=False)
    a3.to_parquet(OUT / "anexo3_clasificacion.parquet", index=False)
    salud.to_parquet(OUT / "salud_componentes.parquet", index=False)

    print(f"Series: {series['fecha'].min():%Y-%m} a {series['fecha'].max():%Y-%m} | "
          f"{series['codigo'].nunique()} series | {len(series):,} filas")
    v = series[series["codigo"].isin(["GEN", "06"])].pivot(index="fecha", columns="codigo", values="var_interanual")
    v = v[["GEN", "06"]].rename(columns={"GEN": "Indice General", "06": "Salud"})
    v = v.loc[[f"{y}-12-01" for y in range(2021, 2026)] + ["2026-03-01"]].round(2)
    print("\nInflacion interanual (%) a fin de ano y marzo 2026:")
    print(v.to_string())
    print("\nSalud por subgrupo (% de Salud) y clasificacion de articulos:")
    print(salud[salud["nivel"] == "subgrupo"][["codigo", "nombre", "ponderacion", "parte_de_salud_pct"]]
          .round(2).to_string(index=False))
    art = salud[salud["nivel"] == "articulo"]
    print(f"Articulos de Salud: {len(art)} | subyacentes: {(art['subyacente'] == 'Subyacente').sum()} | "
          f"bienes: {(art['tipo'] == 'Bien').sum()} | servicios: {(art['tipo'] == 'Servicio').sum()}")


if __name__ == "__main__":
    main()
